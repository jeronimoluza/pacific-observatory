"""INS-Congo -- "Prix moyens mensuels de quelques variétés" retail price
tables, published inside the same monthly "Bulletin INHPC" Excel workbook
as insc_inhpc_cpi.py (see that module's docstring for the shared
discovery path: ins-congo.cg/backend/api/publications/list.php?
sous_secteur=stat-prix, Excel download).

The workbook carries one such table per surveyed city -- Brazzaville,
Pointe-Noire, Dolisie, Owando, Ouesso -- each a genuine official average
RETAIL price series (not an index) for ~30 named food/beverage items plus
3 non-food fuel items (Petrole lampant / kerosene, Charbon de bois /
charcoal, Bois de chauffage / firewood) that are hard-excluded here --
out of scope for this food/beverage-only onboarding pass, same exclusion
pattern as Namibia's zonal food-price fetcher.

Sheets are located by content-sniffing (cell A1 containing both "Prix"
and "variétés") rather than fixed sheet index, since sheet numbering is
not stable across releases -- confirmed live the 5 city tables sit at
sheet positions 6/9/11/13/16 in the July-2026 release, with no fixed
arithmetic relationship between them.

coicop_classification: classifier -- ~30 named grocery items spanning
several COICOP-2018 4-digit classes (cereals, meat, fish, oils,
vegetables, sugar), the same shape Namibia's nsa_zonal_prices.yaml
routes to the classifier rather than a hand-curated map.

Emits PriceObservation rows (analytical_role: official_avg).
"""

from __future__ import annotations

import io
import logging
import re
from datetime import date

import openpyxl
import pandas as pd
from curl_cffi import requests as curl_requests

from prices.fetchers.utils import get_scrape_ts, make_hash

logger = logging.getLogger(__name__)

_PUBLICATIONS_URL = "https://ins-congo.cg/backend/api/publications/list.php"
_COUNTRY = "Congo Republic"
_CURRENCY = "XAF"
_SOURCE_KEY = "cg_insc_retail_prices"
_IDENT = ["source_key", "observation_date", "subnational_area", "item_name"]

# Hard-excluded: non-food/beverage items in the same table (fuel/energy),
# out of scope for this onboarding pass.
_EXCLUDE_ITEMS = {"pétrole lampant", "petrole lampant", "charbon de bois", "bois de chauffage"}

_MONTH_PREFIX = {
    "janv": 1, "fevr": 2, "févr": 2, "mars": 3, "avr": 4, "mai": 5,
    "juin": 6, "juil": 7, "aout": 8, "août": 8, "sept": 9, "oct": 10,
    "nov": 11, "dec": 12, "déc": 12,
}
_MONTH_COL_RE = re.compile(
    r"(janv|f[ée]vr|mars|avr|mai|juin|juil|ao[uû]t|sept|oct|nov|d[ée]c)[a-zé]*-?\s*(\d{2})",
    re.IGNORECASE,
)
_CITY_RE = re.compile(r"variétés\s*à\s*([A-Za-zÀ-ÿ\- ]+)", re.IGNORECASE)


def _clean(v) -> str:
    return re.sub(r"\s+", "", str(v or ""))


def _parse_month_col(label) -> date | None:
    m = _MONTH_COL_RE.search(_clean(label).lower())
    if not m:
        return None
    prefix, yy = m.groups()
    prefix = prefix.replace("û", "u")
    for key, num in _MONTH_PREFIX.items():
        if prefix.startswith(key[:4]) or key.startswith(prefix[:4]):
            return date(2000 + int(yy), num, 1)
    return None


def _find_city_tables(wb) -> list[tuple[object, str]]:
    found: list[tuple[object, str]] = []
    for name in wb.sheetnames:
        ws = wb[name]
        title = str(ws.cell(1, 1).value or "")
        if "prix" in title.lower() and "variétés" in title.lower():
            m = _CITY_RE.search(title)
            city = m.group(1).strip() if m else None
            if city:
                found.append((ws, city))
    return found


def _latest_bulletin(session: curl_requests.Session) -> dict | None:
    resp = session.get(
        _PUBLICATIONS_URL,
        params={"sous_secteur": "stat-prix", "per_page": 100},
        timeout=30,
        impersonate="chrome124",
    )
    resp.raise_for_status()
    pubs = resp.json().get("data", [])
    inhpc = [p for p in pubs if p.get("has_excel")]
    if not inhpc:
        return None
    inhpc.sort(key=lambda p: p.get("id", 0), reverse=True)
    return inhpc[0]


def fetch_cg_insc_retail_prices(cutoff: date) -> pd.DataFrame | None:
    session = curl_requests.Session()

    pub = _latest_bulletin(session)
    if pub is None:
        logger.warning("[%s] No stat-prix bulletin with an Excel export found", _SOURCE_KEY)
        return None

    xlsx_url = "https://ins-congo.cg/" + pub["download_url_excel"].lstrip("/")
    resp = session.get(xlsx_url, timeout=60, impersonate="chrome124")
    if resp.status_code != 200:
        logger.warning("[%s] Excel fetch failed (%s): %s", _SOURCE_KEY, resp.status_code, xlsx_url)
        return None

    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
    city_tables = _find_city_tables(wb)
    if not city_tables:
        logger.warning("[%s] No 'Prix moyens ... variétés' sheets found in %s", _SOURCE_KEY, xlsx_url)
        return None

    rows: list[dict] = []
    for ws, city in city_tables:
        # Header row: first row containing "RUBRIQUE" in column 1.
        header_row = None
        for r in range(1, min(ws.max_row, 10) + 1):
            if "rubrique" in str(ws.cell(r, 1).value or "").lower():
                header_row = r
                break
        if header_row is None:
            continue

        month_cols: list[tuple[int, date]] = []
        for cell in ws[header_row]:
            if cell.column <= 2:  # 1=RUBRIQUE, 2=Unité
                continue
            d = _parse_month_col(cell.value)
            if d is not None:
                month_cols.append((cell.column, d))
        if not month_cols:
            continue

        for r in range(header_row + 1, ws.max_row + 1):
            item_raw = ws.cell(r, 1).value
            unit_raw = ws.cell(r, 2).value
            if item_raw is None:
                continue
            item_name = re.sub(r"\s+", " ", str(item_raw)).strip()
            if not item_name or item_name.lower() in _EXCLUDE_ITEMS:
                continue
            unit = re.sub(r"\s+", "", str(unit_raw or "")) or None

            for col, month_date in month_cols:
                if month_date <= cutoff:
                    continue
                val = ws.cell(r, col).value
                if val is None:
                    continue
                try:
                    price = float(val)
                except (ValueError, TypeError):
                    continue
                if price <= 0:
                    continue
                row = {
                    "observation_date": month_date.isoformat(),
                    "period_kind": "monthly_avg",
                    "country": _COUNTRY,
                    "subnational_area": city,
                    "source_key": _SOURCE_KEY,
                    "item_name": item_name,
                    "price_local": price,
                    "currency": _CURRENCY,
                    "unit": unit,
                    "source_url": xlsx_url,
                    "scrape_ts": get_scrape_ts(),
                    "observation_hash": None,
                }
                row["observation_hash"] = make_hash(row, _IDENT)
                rows.append(row)

    if not rows:
        return None
    return pd.DataFrame(rows)
